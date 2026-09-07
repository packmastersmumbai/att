import sys, io
from docx import Document

p = sys.argv[1]
d = Document(p)

print("=== PARAGRAPHS ===")
for para in d.paragraphs:
    t = para.text.strip()
    if t:
        print(repr(t))

print("\n=== TABLES:", len(d.tables), "===")
for ti, tb in enumerate(d.tables):
    print(f"\n--- table {ti}: {len(tb.rows)} rows x {len(tb.columns)} cols")
    seen = set()
    for ri, row in enumerate(tb.rows):
        cells = []
        for c in row.cells:
            # merged cells repeat the same object; keep it readable
            cells.append(c.text.strip().replace("\n", " / "))
        line = " | ".join(cells)
        if line not in seen:
            print(f"  r{ri}: {line}")
        seen.add(line)
